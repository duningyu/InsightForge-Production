from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
OUTPUT = ROOT / "docs" / "PRD_InsightForge.docx"

NAVY = "111B1A"
INK = "172033"
MUTED = "596662"
BLUE = "3158EB"
BLUE_SOFT = "EAF0FF"
MINT = "148D6F"
MINT_SOFT = "E8F6F1"
AMBER = "B56B08"
AMBER_SOFT = "FFF2DE"
LINE = "D8DFDC"
SURFACE_SOFT = "F5F7F4"
WHITE = "FFFFFF"
RED = "A63D35"
FONT = "Noto Sans CJK SC"
MONO = "Noto Sans Mono CJK SC"


def set_run_font(run, name: str = FONT, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = f"w:{edge}"
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ("val", "sz", "space", "color"):
                if key in edge_data:
                    element.set(qn(f"w:{key}"), str(edge_data[key]))


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph, value: bool = True):
    paragraph.paragraph_format.keep_with_next = value


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8, color=MUTED)


def setup_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 12),
        ("Subtitle", 14, MUTED, 0, 10),
        ("Heading 1", 17, NAVY, 14, 8),
        ("Heading 2", 12.5, BLUE, 11, 5),
        ("Heading 3", 10.5, MINT, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True

    if "IF Small" not in styles:
        small = styles.add_style("IF Small", 1)
        small.font.name = FONT
        small._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        small.font.size = Pt(8)
        small.font.color.rgb = RGBColor.from_string(MUTED)
        small.paragraph_format.space_after = Pt(3)
        small.paragraph_format.line_spacing = 1.1

    if "IF Code" not in styles:
        code = styles.add_style("IF Code", 1)
        code.font.name = MONO
        code._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
        code.font.size = Pt(8)
        code.font.color.rgb = RGBColor.from_string(INK)
        code.paragraph_format.left_indent = Cm(0.45)
        code.paragraph_format.right_indent = Cm(0.45)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(5)


def setup_sections(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.55)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("INSIGHTFORGE 2.0  ·  GUIDED EVIDENCE WORKSPACE")
    set_run_font(run, size=7.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    table = footer.add_table(rows=1, cols=3, width=Inches(6.8))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, center, right = table.rows[0].cells
    for cell in (left, center, right):
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell, top={"val": "single", "sz": "4", "color": LINE})
    p = left.paragraphs[0]
    set_run_font(p.add_run("Implementation-aligned PRD"), size=7.5, color=MUTED)
    p = center.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("2026-08-26"), size=7.5, color=MUTED)
    add_page_number(right.paragraphs[0])
    if not fp.text:
        fp._element.getparent().remove(fp._element)


def add_cover(doc: Document):
    for _ in range(2):
        doc.add_paragraph("")
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = False
    cell = band.cell(0, 0)
    cell.width = Cm(17)
    shade_cell(cell, NAVY)
    set_cell_margins(cell, 420, 400, 420, 400)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("INSIGHTFORGE 2.0")
    set_run_font(r, size=31, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    r = p2.add_run("Guided Evidence Workspace")
    set_run_font(r, size=19, bold=True, color="A7C0FF")
    p3 = cell.add_paragraph()
    r = p3.add_run("完整产品需求与策略文档")
    set_run_font(r, size=13, color="DCE5E1")

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("从“技术型证据工作台”升级为“真正帮助小白完成产品项目的可追溯产品教练”。")
    set_run_font(r, size=16, bold=True, color=INK)
    p = doc.add_paragraph()
    r = p.add_run("默认：对话式引导 + 分步任务流 + 右侧证据追溯面板｜可切换：高级工作台")
    set_run_font(r, size=11.5, color=BLUE)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    metadata = [
        ("文档版本", "v2.0.0"),
        ("文档状态", "Implementation-aligned Release PRD"),
        ("日期", "2026-08-26"),
        ("本轮反馈证据", "一次定性产品经理反馈 + v1.1.0 代码与交互审计"),
        ("主张边界", "本轮反馈不构成市场验证；参数、Agent 与业务价值均按证据边界表述"),
    ]
    for i, (key, value) in enumerate(metadata):
        c1, c2 = table.rows[i].cells
        c1.width = Cm(4.1)
        c2.width = Cm(12.2)
        shade_cell(c1, MINT_SOFT if i % 2 == 0 else BLUE_SOFT)
        shade_cell(c2, WHITE)
        for c in (c1, c2):
            set_cell_margins(c, 100, 120, 100, 120)
            set_cell_border(c, top={"val": "single", "sz": "4", "color": LINE}, bottom={"val": "single", "sz": "4", "color": LINE}, left={"val": "single", "sz": "4", "color": LINE}, right={"val": "single", "sz": "4", "color": LINE})
        set_run_font(c1.paragraphs[0].add_run(key), size=8.8, bold=True, color=MUTED)
        set_run_font(c2.paragraphs[0].add_run(value), size=8.8, color=INK)

    doc.add_paragraph("")
    add_callout(doc, "可信度声明", "本 PRD 将用户反馈、已实现代码、系统建议和待验证假设分开记录。任何来自模拟材料、模型输出或单次定性反馈的结论，都不得升级为“市场已经验证”。", AMBER_SOFT, AMBER)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.page_break_before = True
    return p


def add_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None, small: bool = False):
    p = doc.add_paragraph(style="IF Small" if small else "Normal")
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True, color=INK, size=8 if small else 9.4)
        set_run_font(p.add_run(text[len(bold_prefix):]), color=MUTED if small else INK, size=8 if small else 9.4)
    else:
        set_run_font(p.add_run(text), color=MUTED if small else INK, size=8 if small else 9.4)
    return p


def add_bullets(doc: Document, items: Iterable[str], level: int = 0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
        set_run_font(p.add_run(item), size=9, color=INK)


def add_numbered(doc: Document, items: Iterable[str]):
    """Add an independently numbered list that restarts at 1 for every call.

    Word/LibreOffice may continue the built-in ``List Number`` sequence across
    unrelated sections. Manual prefixes keep the rendered PRD deterministic.
    """
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.62)
        p.paragraph_format.first_line_indent = Cm(-0.48)
        set_run_font(p.add_run(f"{index}. "), size=9, bold=True, color=INK)
        set_run_font(p.add_run(item), size=9, color=INK)


def add_callout(doc: Document, title: str, body: str, fill: str = BLUE_SOFT, accent: str = BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, 130, 170, 130, 170)
    set_cell_border(cell, left={"val": "single", "sz": "20", "color": accent})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(title), size=9.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    set_run_font(p2.add_run(body), size=8.8, color=INK)
    return table


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None, font_size: float = 7.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Cm(widths[i])
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 95, 95, 95, 95)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(str(header)), size=font_size, bold=True, color=WHITE)
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            shade_cell(cell, WHITE if row_index % 2 == 0 else SURFACE_SOFT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, 90, 95, 90, 95)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(str(value)), size=font_size, color=INK)
    doc.add_paragraph("").paragraph_format.space_after = Pt(0)
    return table


def add_picture(doc: Document, path: Path, width_cm: float, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(path), width=Cm(width_cm))
    # python-docx does not expose alt text directly, so set the drawing
    # properties used by Word and accessibility readers.
    shape._inline.docPr.set("title", path.stem.replace("_", " "))
    shape._inline.docPr.set("descr", caption)
    cp = doc.add_paragraph(style="IF Small")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cp.add_run(caption), size=7.8, color=MUTED)


def add_code(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "EEF1EF")
    set_cell_margins(cell, 110, 140, 110, 140)
    p = cell.paragraphs[0]
    p.style = doc.styles["IF Code"]
    set_run_font(p.add_run(text), name=MONO, size=7.8, color=INK)


def build_document() -> Document:
    doc = Document()
    setup_styles(doc)
    setup_sections(doc)
    add_cover(doc)

    add_heading(doc, "修订记录", 1)
    add_table(doc, ["版本", "日期", "核心变化", "状态"], [
        ["v1.0", "2026-08-24", "完整产品策略与 Evidence Workspace 规划：用户、竞品、来源治理、项目级 RAG、PRD/TechDoc、质量 Loop、MCP。", "历史基线"],
        ["v1.1.0", "2026-08-26", "形成可运行技术工作台；完成 Canvas、来源、RAG、文档、验证、审批、导出、Function Calling 与 MCP，但信息架构仍偏开发者。", "被 v2.0.0 替代"],
        ["v2.0.0", "2026-08-26", "默认新手引导；对话 + 分步任务 + 右侧证据面板；检索运行可回放；Claim-Evidence Ledger；真实 Handoff ZIP；高级工作台保留。", "当前实现"],
    ], widths=[2.2, 2.6, 9.2, 2.4], font_size=7.5)

    add_heading(doc, "文档导航", 1)
    nav = [
        "01 执行摘要与明确产品判断",
        "02 用户反馈与版本迭代",
        "03 项目背景、目标与非目标",
        "04 用户分层、Persona 与 JTBD",
        "05 竞品类别与差异化战略",
        "06 产品定位、价值主张与设计原则",
        "07 用户旅程、主流程与异常流程",
        "08 信息架构、双模式 UI 与响应式设计",
        "09 需求池、优先级与版本范围",
        "10 核心功能详细需求与验收",
        "11 受限 PM Coach、Agent Harness、MCP 与 LangChain 边界",
        "12 RAG 配置、检索轨迹与来源治理",
        "13 Claim-Evidence Ledger、文档生成与质量 Gate",
        "14 数据模型、API 与审计",
        "15 真实 AI Coding 交接 ZIP",
        "16 指标体系、RAG Evals 与实验协议",
        "17 用户测试、上线策略与运营埋点",
        "18 风险、失败退出条件与主张边界",
        "19 Roadmap 与当前实现审计",
    ]
    add_table(doc, ["章节", "内容"], [[item[:2], item[3:]] for item in nav], widths=[2.0, 14.3], font_size=8)

    add_heading(doc, "1. 执行摘要与明确产品判断", 1)
    add_callout(doc, "明确结论", "InsightForge 2.0 不再是“让用户把专业字段填完后帮他生成 PRD”的技术后台，而是面向小白、转岗者和早期产品实践者的证据可追溯产品教练。核心交互采用“对话式引导 + 分步任务流 + 右侧证据追溯面板”，且新手引导模式为默认，右上角可切换高级工作台。", MINT_SOFT, MINT)
    add_paragraph(doc, "产品存在理由不是文本生成速度，而是让用户在缺少成熟产品方法与组织上下文的情况下，仍能完成问题澄清、证据补充、方案比较、版本化需求合同、文档审查和可执行交接。")
    add_table(doc, ["维度", "v1.1.0 状态", "v2.0.0 产品判断"], [
        ["用户入口", "总览 / Canvas / 来源库 / RAG / Loop / MCP", "按用户任务进入：想法 → 用户问题 → 证据 → 方案 → 文档 → 交接"],
        ["小白负担", "用户先理解专业字段和参数", "系统一次只问一个问题，说明原因、给示例，并允许“不确定”"],
        ["可信度", "有 source/chunk ID，但生成和检索过程难回放", "来源元数据 + retrieval run + claim link + approval + manifest 全链记录"],
        ["RAG", "Top-K 和权重可见但依据不清", "命名档位、用途解释、trade-off、运行记录；不得声称参数最优"],
        ["AI Coding", "MCP Prompt/接口描述", "真实 AI Coding 交接 ZIP：批准版、来源、主张、检索、任务、测试与 SHA-256"],
    ], widths=[2.5, 6.3, 7.7])
    add_paragraph(doc, "北极星不定义为“生成多少字”，而定义为 Evidence-backed Project Completion Rate：用户是否完成关键项目阶段，且事实、假设、建议和未知项被正确区分并可追溯。")

    add_heading(doc, "2. 用户反馈与版本迭代", 1)
    add_callout(doc, "反馈证据边界", "本轮输入来自一次定性产品经理反馈，用于发现产品缺陷与设计方向；它不构成市场验证，不能据此主张需求规模、留存、付费意愿或商业价值已经成立。", AMBER_SOFT, AMBER)
    add_heading(doc, "2.1 反馈内容", 2)
    feedback_rows = [
        ["UI 与信息架构", "界面仍像技术后台，需要更好看；导航以技术模块而不是用户任务组织。", "重构为双模式：Guided Mode 默认；Advanced Workspace 明确切换。"],
        ["新手澄清", "约束不知道怎么填，不知道想法与澄清应该写到什么粒度。", "PM Coach 一次一个问题，提供 why / examples / choices；系统建议不自动保存。"],
        ["来源库", "不知道证据来源库该怎么写，也不知道 authority 数字代表什么。", "普通语言来源类别；展示能/不能证明什么；保存 authority basis、URL、发布者、时间和 SHA。"],
        ["RAG 配置", "不知道配置是干什么的，不知道 Top-K 与权重有什么用。", "快速/平衡/充分查证档位；页面回答“为什么使用当前检索档位”。"],
        ["参数依据", "返回数量等没有可追溯来源。", "每个 retrieval run 记录 profile、Top-K、weights、候选数、命中和 actor。"],
        ["黑箱", "看不到来源、检索和主张链路，无法判断系统是否在说谎。", "右侧 Evidence Panel + Claim-Evidence Ledger + 原始 URL + run 回放。"],
        ["AI Coding", "交接是虚无的，只有名义接口。", "真实 AI Coding 交接 ZIP，绑定批准文档并包含 manifest 与验收任务。"],
        ["Agent/MCP/LangChain", "希望智能体主动给予必要回答。", "实现受限 PM Coach 与 MCP 资源/工具；不为技术标签强行引入 LangChain。"],
    ]
    add_table(doc, ["主题", "反馈/问题", "v2.0.0 回应"], feedback_rows, widths=[2.5, 6.0, 8.0], font_size=7.4)
    add_heading(doc, "2.2 根因", 2)
    add_bullets(doc, [
        "目标用户是初学者，但界面默认假设用户已经理解 PRD、产品约束、来源类型、authority、Top-K、评测和交接。",
        "系统把“如何做产品”的核心认知工作重新交给用户，仅把结果结构化。",
        "引用格式合法与引用语义支持没有充分区分；旧生成器曾按索引轮换引用，形成表面可信感。",
        "AI Coding Handoff 没有形成可下载、可校验、绑定批准版本的产物。",
        "技术名词成为信息架构，而不是隐藏在受控执行链中的实现细节。",
    ])
    add_heading(doc, "2.3 版本迭代原则", 2)
    add_numbered(doc, [
        "先降低认知负担，再增加 Agent 能力；不能用自由聊天掩盖流程缺陷。",
        "先建立可追溯链，再追求生成丰富度；Evidence > Fluency。",
        "模型建议、来源事实、用户确认和待验证项必须是不同状态。",
        "高级能力保留，但默认不暴露给小白。",
        "新增模块必须有独立验收与失败退出条件；不因简历技术标签而保留。",
    ])

    add_heading(doc, "3. 项目背景、目标与非目标", 1)
    add_heading(doc, "3.1 背景", 2)
    add_paragraph(doc, "AI 产品求职者、转岗者和独立开发者可以用大模型快速得到一份流畅 PRD，但项目仍常失败在：问题未澄清、资料来源混乱、事实与假设混淆、文档版本漂移、实现交接缺少边界。")
    add_table(doc, ["现象", "根因", "v2 产品机会"], [
        ["只有模糊想法，不知道先做什么", "缺少可理解的 discovery 顺序", "用一问一答和六步任务流降低启动门槛"],
        ["不知道约束和指标怎么写", "模板告诉用户填什么，却不解释为什么", "Coach 提供场景化示例、选择和建议草案"],
        ["AI 输出看似完整但无法验证", "来源、检索和主张没有映射", "Source → Retrieval Run → Claim → Document 的证据链"],
        ["交给 AI Coding 后偏离需求", "上下文、版本、非目标和验收不稳定", "批准版 Handoff ZIP + manifest + AGENTS.md"],
    ], widths=[4.6, 5.4, 6.5])
    add_heading(doc, "3.2 产品目标", 2)
    add_bullets(doc, [
        "让无产品经验用户只用自然语言即可完成可解释的项目定义。",
        "让每条重要主张可追溯到用户确认、来源证据或明确的待验证状态。",
        "让 RAG 运行配置、候选与命中可回放，而不是只返回结果。",
        "让 PRD/TechDoc 经过确定性检查和人工 Gate，再进入交接。",
        "让 AI Coding 读取批准、版本化、带哈希的上下文，而不是聊天记忆。",
    ])
    add_heading(doc, "3.3 非目标", 2)
    add_bullets(doc, [
        "不宣称替代资深产品经理或自动完成完整 discovery。",
        "不宣称引用等于事实正确；系统只提高可追溯性和可审计性。",
        "不在 v2 强制接入 LangChain、GraphRAG、远程 MCP OAuth 或多 Agent 群聊。",
        "不实现企业级实时协作、Portfolio Roadmap、SSO 或多租户权限。",
        "不使用一次定性反馈主张市场验证、留存或付费价值。",
    ])

    add_heading(doc, "4. 用户分层、Persona 与 JTBD", 1)
    add_table(doc, ["用户", "典型状态", "核心任务", "主要风险", "优先级"], [
        ["AI 产品求职者", "需要做 1–2 个能解释的作品集项目", "从想法到证据、PRD、技术边界和可运行交付", "为了包装而虚构用户研究/指标", "核心"],
        ["转岗/初级 PM", "了解基础术语，但不熟悉完整 discovery 与 AI 评测", "获得方法引导、证据判断和下一步建议", "被专业字段和参数劝退", "核心"],
        ["产品能力较弱的独立开发者", "有 AI Coding 能力，产品定义弱", "明确用户、范围、验收和交接", "直接编码导致范围漂移", "次核心"],
        ["成熟企业产品团队", "已有 Productboard/Notion/Miro/内部知识库", "组织级协作、权限和反馈闭环", "InsightForge 功能过轻", "非当前核心"],
    ], widths=[3.1, 4.2, 4.4, 3.5, 1.4])
    add_heading(doc, "4.1 Persona", 2)
    add_table(doc, ["Persona", "目标", "现在怎么做", "成功体验"], [
        ["小宁｜转岗 AI PM", "完成一个可面试讲解的 AI 产品项目", "问大模型、复制模板、到处存资料", "系统告诉他下一步、为什么、证据够不够，并能导出完整交接"],
        ["阿杰｜独立开发者", "在编码前锁定问题与验收", "聊天里讨论后直接让 Codex 开发", "批准版上下文和任务可执行，返工与需求漂移减少"],
        ["小林｜初级 PM", "理解 RAG/Agent 产品设计而非只写术语", "看文档但不知道参数和评测关系", "能查看每次检索配置、trade-off、主张支持关系和失败案例"],
    ], widths=[3.2, 4.3, 4.4, 5.3])
    add_heading(doc, "4.2 JTBD", 2)
    add_table(doc, ["When", "I want to", "So I can"], [
        ["当我只有一个模糊想法", "有人用我能回答的问题一步步澄清", "避免一开始面对空白 PRD 和专业字段"],
        ["当我不知道约束/指标怎么写", "看到为什么要问、示例和常见选择", "做出真实而不是“专业腔”的项目合同"],
        ["当我加入资料", "知道它可以/不能证明什么", "不把模拟材料、模型输出或实现证据写成用户事实"],
        ["当系统给出建议", "看到来源、查询、参数和未知项", "判断建议是否可信并能回到原始材料"],
        ["当我要交给 AI Coding", "导出批准版本和验收边界", "避免工具读取旧版本或自由补需求"],
    ], widths=[5.0, 5.8, 6.0])

    add_heading(doc, "5. 竞品类别与差异化战略", 1)
    add_callout(doc, "研究范围", "本轮主要目标是基于 v1.0 文档与当前实现完成版本迭代，不重新声称最新竞品能力。原 v1.0 中 ChatPRD、Productboard Spark、Notion AI、Miro AI 等公开能力判断应在正式外部发布前再次核验。", BLUE_SOFT, BLUE)
    add_table(doc, ["竞品类别", "优势", "不应硬拼", "InsightForge 差异"], [
        ["Chat-first PRD 助手", "上手快、生成流畅、模板成熟", "生成速度和文案长度", "一问一答仍绑定 Canvas、证据、运行和人工 Gate"],
        ["企业 Product Intelligence", "客户反馈、Roadmap、权限与协作强", "组织级数据规模和协作生态", "面向个人早期实践者，强调本地、来源边界和作品集可信度"],
        ["通用知识工作台", "连接器、搜索、协作和知识沉淀", "通用工作空间能力", "专门定义产品项目阶段、主张类型和交接 Gate"],
        ["原型/AI Coding Agent", "实现速度快、可直接改代码", "代码生成和 IDE 集成数量", "先锁定批准上下文、验收、非目标和证据再编码"],
    ], widths=[3.3, 4.3, 4.0, 5.2])
    add_heading(doc, "5.1 Why InsightForge", 2)
    add_bullets(doc, [
        "Novice-first：用户不需要先学会 PRD 才能使用产品。",
        "Evidence-first：系统不仅生成答案，还展示来源、检索和主张关系。",
        "Version-first：Canvas、文档和 Handoff 都绑定版本，不依赖聊天记忆。",
        "Handoff-first：项目终点不是一份文档，而是可执行且可校验的交付包。",
        "Claim-boundary-first：明确“可以怎么表述”和“不能怎么表述”。",
    ])

    add_heading(doc, "6. 产品定位、价值主张与设计原则", 1)
    add_callout(doc, "Positioning Statement", "For 想把模糊想法做成完整 AI 产品项目、但缺少成熟产品流程和组织上下文的早期产品实践者，InsightForge 2.0 是一个 novice-first, evidence-backed product project coach。它通过分步引导、来源治理、可回放 RAG、主张账本与批准版交接，帮助用户完成可解释、可审查、可开发的项目。", MINT_SOFT, MINT)
    add_table(doc, ["原则", "解释", "产品行为"], [
        ["Task > Module", "用户先理解当前任务，不需要理解内部系统模块", "默认六步任务流；技术页面仅在 Advanced Workspace"],
        ["Project > Document", "项目完成包含问题、证据、决策、文档和交接", "进度不是“PRD 已生成”就结束"],
        ["Evidence > Fluency", "语言流畅不等于事实可靠", "重要主张显示来源/状态/未知项"],
        ["Confirmation > Assumption", "系统建议不能静默成为需求", "proposed_canvas_patch 需要显式 apply"],
        ["Trace > Trust me", "可信度来自可回放链路", "检索 Run、Claim Link、Audit、Manifest"],
        ["Bounded > Autonomous", "高风险写入必须停在人工 Gate", "Agent 不能审批、发布、删除或覆盖 approved"],
    ], widths=[3.0, 6.0, 7.5])

    add_heading(doc, "7. 用户旅程、主流程与异常流程", 1)
    add_picture(doc, ASSETS / "user_flow.png", 16.8, "图 1｜v2.0 新手主流程：从一句话想法到真实 AI Coding 交接")
    add_heading(doc, "7.1 主流程", 2)
    add_numbered(doc, [
        "创建项目：只要求名称和一句话想法，不要求先完成 Canvas。",
        "产品教练澄清：按 audience → problem → constraints → evidence → success 逐步推进。",
        "方案比较：提供 3 个 alternatives，并展示收益、成本、风险和未知项。",
        "确认项目合同：系统生成 proposed Canvas；用户显式确认后保存新版本。",
        "补充证据：用户用普通语言选择来源类别；系统解释允许主张与局限。",
        "检索与查证：运行命名 Profile，保存 Run ID、配置、候选和结果。",
        "生成与审查：创建 PRD/TechDoc；查看 Claim-Evidence Ledger；执行 Validator。",
        "人工审批：只有 passed 版本才允许 approve；审批人和 note 进入 Audit。",
        "交接：PRD+TechDoc 均批准后导出 ZIP，记录 run ID 和 SHA-256。",
    ])
    add_heading(doc, "7.2 异常流程", 2)
    add_table(doc, ["异常", "系统行为", "用户恢复路径"], [
        ["用户不知道怎么回答", "展示示例、选择和“不确定/帮我分析”；不编造用户事实", "填写最小已知信息，未知项进入 evidence gap"],
        ["来源类型不确定", "提供 provisional classification、basis 和 needs_confirmation", "用户在 Advanced Workspace 复核/修正"],
        ["检索无结果", "保存空 Run；不允许 authority-only 召回", "换查询、增加来源或改用充分查证档位"],
        ["来源冲突", "保留支持/反驳/上下文关系，不让模型自行裁决", "人工选择是否形成决策或标 unresolved"],
        ["文档验证失败", "最多两轮定向修复，仍失败则 needs_human_review", "补充证据、修改 Canvas 或手工编辑新版本"],
        ["Handoff 阻断", "缺 PRD/TechDoc approval 时 fail closed", "生成、验证、阅读并批准缺失文档"],
    ], widths=[4.0, 6.2, 6.3])

    add_heading(doc, "8. 信息架构、双模式 UI 与响应式设计", 1)
    add_picture(doc, ASSETS / "architecture_modes.png", 16.8, "图 2｜Guided Mode 与 Advanced Workspace 共用同一控制面")
    add_heading(doc, "8.1 Guided Mode", 2)
    add_table(doc, ["区域", "作用", "关键组件"], [
        ["左侧任务轨", "让用户知道当前在哪一步、还差什么", "6 步、complete/current/locked 状态、Canvas 版本"],
        ["中间 Coach", "一次一个问题并形成可回放记录", "question、why、examples、choices、conversation、apply Gate"],
        ["右侧 Evidence Panel", "让用户在同一屏看到证据而非相信黑箱", "来源、检索轨迹、主张；原始链接与 SHA"],
    ], widths=[3.5, 6.1, 6.8])
    add_picture(doc, ASSETS / "ui_guided.png", 14.8, "图 3｜Guided Mode 本地静态验证截图（演示数据，不是在线用户数据）")
    add_heading(doc, "8.2 Advanced Workspace", 2)
    add_paragraph(doc, "右上角切换高级工作台。它不是另一个产品，而是同一项目数据的专业视图，用于查看原始字段、RAG 参数、文档版本、主张、审计和 Handoff。")
    add_picture(doc, ASSETS / "ui_advanced.png", 14.8, "图 4｜Advanced Workspace 本地静态验证截图（演示数据）")
    add_heading(doc, "8.3 响应式与可用性", 2)
    add_bullets(doc, [
        "桌面：三列布局，左/右侧可保持 sticky，主对话位于视觉中心。",
        "平板：侧栏与证据面板转为单列或横向滚动，不隐藏核心状态。",
        "移动端：模式标签仍可见；任务轨、Coach、Evidence 顺序排列。",
        "支持 keyboard focus、aria-selected/aria-current 和 prefers-reduced-motion。",
        "中文主任务语言优先；内部英文 ID 只在 Advanced/审计视图出现。",
    ])

    add_heading(doc, "9. 需求池、优先级与版本范围", 1)
    add_table(doc, ["需求", "用户价值", "v2 状态", "优先级", "成功 Gate"], [
        ["Guided Mode 默认", "降低启动与理解成本", "已实现", "P0", "小白无需培训完成首轮澄清"],
        ["Coach 状态机", "一次一个问题、建议与确认分离", "已实现", "P0", "无静默写入；可回放"],
        ["来源引导与 provenance", "知道资料可支持什么", "已实现", "P0", "分类/边界理解率达标"],
        ["Retrieval Trace", "知道系统搜了什么和为什么", "已实现", "P0", "每次检索都有 run/config/hits"],
        ["Claim-Evidence Ledger", "判断一句话是否有证据", "已实现", "P0", "source-backed claim 有合法 link"],
        ["Human Approval", "区分草稿和批准版", "已实现", "P0", "passed + explicit human confirm"],
        ["真实 Handoff ZIP", "把稳定上下文交给 AI Coding", "已实现", "P0", "批准 PRD+TechDoc + valid manifest"],
        ["外部网页自动抓取", "降低来源录入成本", "未实现", "P1", "先解决安全、robots、内容哈希与失效"],
        ["冻结 RAG Eval 控制台", "让参数选择有验证依据", "未实现", "P1", "validation-only 选择 + frozen report"],
        ["远程 MCP/团队协作", "多人、企业集成", "未实现", "P2", "真实企业需求和权限模型成立"],
    ], widths=[3.3, 4.6, 2.4, 1.5, 4.7], font_size=7.2)

    add_heading(doc, "10. 核心功能详细需求与验收", 1)
    features = [
        ("F1 Project & Guided PM Coach", [
            ["用户价值", "用自然语言完成项目澄清，不需要先理解 PRD 字段。"],
            ["主流程", "项目 summary 进入 idea；Coach 依次收集 audience/problem/constraints/evidence/success；solution 阶段输出 3 个 alternatives；canvas_ready 需要 apply。"],
            ["状态语义", "confirmed_fields 与 suggested_fields 分开；messages 记录 role、type、metadata 和时间。"],
            ["验收", "每轮只出现一个核心问题；why/examples/choices 非空；未 apply 前 project_canvas 不变化；apply 后 version+1。"],
            ["边界", "Coach 不能审批文档、发布、删除项目或覆盖批准版本。"],
        ]),
        ("F2 Source Guidance & Provenance", [
            ["用户价值", "不用理解枚举和 0–1 authority，也能正确添加证据。"],
            ["主流程", "选择真实访谈/官网/公开报告/自己的输入/模拟材料/模型输出/实现证据/不确定；系统映射内部类型。"],
            ["解释", "显示 allowed_claims、limitations、recommended authority、authority basis、needs_confirmation。"],
            ["持久化", "URL、publisher、published_at、captured_at、authority_label/basis、status、metadata、SHA-256。"],
            ["验收", "unknown 分类为 provisional；模拟材料不能自动变成 real_user_research；来源可回到 URL 或明确无链接。"],
        ]),
        ("F3 Retrieval Profiles & Trace", [
            ["用户价值", "理解检索档位和结果来源，不再面对无法解释的 Top-K。"],
            ["主流程", "选择 quick/balanced/thorough；高级模式可显式 override Top-K；系统先 project filter，再 rank。"],
            ["Run 字段", "query、purpose、profile、Top-K source、weights、filters、candidate/returned、actor、time、hits。"],
            ["验收", "同一 run 可通过 ID 回放；空结果仍保存；跨项目 chunk 永不进入候选。"],
            ["边界", "当前配置标 manual_baseline_not_frozen_best；不得声称参数最优。"],
        ]),
        ("F4 Document Generation & Claim Ledger", [
            ["用户价值", "生成 PRD/TechDoc 的同时知道每个关键主张属于哪种证据状态。"],
            ["主流程", "DocumentLoop 构建 Evidence Package，生成 content/citations/claims，执行最多两轮 Validator，保存不可变版本。"],
            ["对齐", "公开竞品陈述优先 public_source；实现陈述使用 implementation_evidence；Canvas 内容为 user_confirmed；不足则 unresolved。"],
            ["验收", "不再按索引轮换引用；source_backed claim 必须有合法 chunk link；模拟研究在局部披露。"],
            ["边界", "LLM 输出若无法语义验证，保守标 unresolved；系统不宣称自动真值验证。"],
        ]),
        ("F5 Validation, Approval & Export", [
            ["用户价值", "明确草稿、可审查版本和批准版本。"],
            ["验证", "必要章节、引用存在/作用域、来源披露、Canvas 一致性、claim link/type。"],
            ["审批", "只有 validation_status=passed 且 human_confirmed=true；记录 actor/note/time。"],
            ["导出", "Markdown、JSON、DOCX；批准版本不可覆盖。"],
            ["验收", "L2 工具未确认时 403；失败版本可读但不可 approve。"],
        ]),
        ("F6 AI Coding Handoff", [
            ["用户价值", "将可执行、稳定、带边界的上下文交给 Codex/Claude Code/Cursor/通用 Agent。"],
            ["Readiness", "Canvas + approved/passed PRD + approved/passed TechDoc；缺一则 BLOCKED。"],
            ["产物", "批准文档、Context、Sources、Claims、Retrieval Trace、Acceptance Tests、Tasks、AGENTS.md、Manifest。"],
            ["验收", "每个文件有 SHA-256；package SHA-256 与 Header 返回；unresolved claim 进入包。"],
            ["边界", "MCP 仅提供 readiness/manifest preview；二进制 ZIP 仍需显式 UI/HTTP 操作。"],
        ]),
        ("F7 Advanced Workspace & Audit", [
            ["用户价值", "为专业用户提供原始参数、版本、工具和审计，而不打扰小白主流程。"],
            ["模块", "Overview、Canvas、Sources、Retrieval、Documents、Handoff、Audit & Tools。"],
            ["审计", "actor/action/entity/payload/time；ToolRegistry 记录 risk 和 arguments。"],
            ["验收", "Guided/Advanced 共享项目数据；切换不复制状态；移动端仍显示模式标签。"],
            ["边界", "高级页面可见不等于用户必须配置；默认仍使用安全 profile。"],
        ]),
    ]
    for title, rows in features:
        add_heading(doc, title, 2)
        add_table(doc, ["项", "要求"], rows, widths=[3.0, 13.5], font_size=7.7)

    add_heading(doc, "11. 受限 PM Coach、Agent Harness、MCP 与 LangChain 边界", 1)
    add_heading(doc, "11.1 为什么不是自由聊天 Agent", 2)
    add_paragraph(doc, "把一个自由聊天框放进 v1.1.0 会把表单黑箱变成聊天黑箱。v2 采用 bounded agent harness：状态、工具、写入、风险和中止条件由应用控制。")
    add_code(doc, "User answer → GuidedProjectService state transition → proposed patch\n            → explicit apply → ProjectService → Canvas version\n\nOptional LLM wording → same ToolRegistry → same validation / approval gates")
    add_table(doc, ["层", "职责", "不能做"], [
        ["PM Coach", "澄清、解释、给示例/alternatives、形成建议草案", "把建议当成用户事实；自动审批/发布"],
        ["Domain Services", "项目、来源、检索、文档、主张、交接的确定性语义", "绕过 project scope 或 immutable version"],
        ["ToolRegistry", "严格 JSON Schema、L0/L1/L2 风险和审计", "注册删除、覆盖 approved、外部发布"],
        ["MCP", "向外部客户端暴露项目资源与受限工具", "判断事实真伪或自动授权"],
        ["Optional LLM", "改善表达、理解自然语言、建议查询", "改变授权/验证/写入语义"],
    ], widths=[3.0, 7.0, 6.5])
    add_heading(doc, "11.2 LangChain 决策", 2)
    add_callout(doc, "明确技术判断", "LangChain 不作为本版本强制依赖。当前 FastAPI + Domain Services + ToolRegistry + SQLite 已满足状态、工具、审计和人工 Gate。只有真实出现多模型、多 MCP Server、可暂停图状态、复杂分支恢复和共享 middleware 需求时，才评估 LangGraph/LangChain。", BLUE_SOFT, BLUE)
    add_paragraph(doc, "不引入 LangChain 的原因不是否定框架，而是防止把“增加依赖”误写成“解决黑箱”。来源追溯、Claim-Evidence、运行记录和权限 Gate 必须由产品数据契约实现。")
    add_heading(doc, "11.3 MCP 产品边界", 2)
    add_bullets(doc, [
        "Resources：Canvas、来源、文档版本、Claim Ledger、Handoff readiness。",
        "Tools：project-scoped retrieval、创建草稿、读取 claims/readiness、准备 manifest preview。",
        "Approval 仍是 L2；host 必须传入 human_confirmed。",
        "Binary Handoff export 不开放给模型自动执行。",
        "当前为本地 STDIO；不能表述为远程生产 OAuth/企业权限。",
    ])

    add_heading(doc, "12. RAG 配置、检索轨迹与来源治理", 1)
    add_heading(doc, "12.1 Profile 设计", 2)
    add_table(doc, ["Profile", "Top-K", "用户解释", "Trade-off", "状态"], [
        ["quick_explore_v1", "4", "快速了解资料方向", "阅读成本低，可能漏弱相关证据", "manual baseline"],
        ["balanced_traceable_v1", "8", "平衡查证（默认）", "召回和人工负担折中", "manual baseline"],
        ["thorough_review_v1", "12", "充分查证", "覆盖更广，噪声和复核成本更高", "manual baseline"],
        ["document_generation_v1", "5/query", "内部 Evidence Package", "不是面向用户的全局最佳配置", "manual baseline"],
    ], widths=[4.0, 2.0, 4.2, 4.2, 2.2], font_size=7.5)
    add_code(doc, "hybrid = 0.55 * normalized_BM25\n       + 0.30 * normalized_TF_IDF_cosine\n       + 0.15 * authority")
    add_callout(doc, "为什么使用当前检索档位", "界面必须同时展示：它控制什么、为什么使用当前值、调大/调小的代价、选择依据、validation status 和 run ID。当前 selection basis 是工程基线，标记为 manual_baseline_not_frozen_best。", MINT_SOFT, MINT)
    add_heading(doc, "12.2 可回放字段", 2)
    add_table(doc, ["对象", "字段", "用途"], [
        ["retrieval_runs", "query, purpose, profile_id, top_k, weights, filters, candidate_count, returned_count, actor, created_at", "重现“系统搜了什么、用什么设置、为什么”"],
        ["retrieval_hits", "rank, source_id, chunk_id, component scores, hybrid_score", "解释每个结果为何入选"],
        ["sources", "url, publisher, published/captured time, authority basis, status, SHA", "回到原始来源并判断时效/可信度"],
    ], widths=[3.0, 8.3, 5.2])
    add_picture(doc, ASSETS / "provenance_chain.png", 16.8, "图 5｜Source → Retrieval Run → Claim → Document/Handoff 的可追溯链")
    add_heading(doc, "12.3 来源规则", 2)
    add_table(doc, ["内部类型", "允许主张", "禁止升级"], [
        ["real_user_research", "在样本范围内作为真实用户证据", "不能概括为整体市场结论"],
        ["simulated_research", "演示可能场景和测试样本", "不能写成真实用户调研"],
        ["public_source", "公开页面/报告实际写明的事实", "不能推断真实用户价值或效果"],
        ["user_input", "项目发起人的目标、选择和约束", "不能自动代表外部用户需求"],
        ["model_hypothesis", "待验证建议或假设", "不能变成已验证事实"],
        ["implementation_evidence", "代码、测试、API、运行状态", "不能证明业务价值、留存或市场效果"],
    ], widths=[4.0, 6.4, 6.1])

    add_heading(doc, "13. Claim-Evidence Ledger、文档生成与质量 Gate", 1)
    add_heading(doc, "13.1 Claim 类型", 2)
    add_table(doc, ["Claim Type", "定义", "引用要求", "UI 表示"], [
        ["user_confirmed", "来自 Canvas/用户明确确认的项目合同", "可不引用外部来源，但需记录 Canvas version", "用户已确认"],
        ["source_backed", "由项目内来源支持的事实性主张", "必须绑定 source/chunk；记录 retrieval run", "来源支持"],
        ["model_suggestion", "系统建议、方案或风险提示", "必须局部标“建议/待验证”", "系统建议"],
        ["unresolved", "证据不足、冲突或 LLM 输出未完成语义验证", "不允许包装成确定事实", "待补证"],
    ], widths=[3.6, 6.0, 4.2, 2.7])
    add_heading(doc, "13.2 修复旧引用问题", 2)
    add_paragraph(doc, "v1.1.0 的本地生成器曾按 Evidence 列表索引轮换引用。引用 ID 虽然合法，却可能与当前句子的来源类型不匹配。v2 通过 source-type + lexical fit 选择证据，并持久化 Claim 与 Evidence link。")
    add_table(doc, ["主张类别", "优先证据", "不足时行为"], [
        ["竞品公开能力/市场公开事实", "public_source", "标 unresolved，不用 simulated/implementation 填充"],
        ["当前系统实现", "implementation_evidence", "仅陈述代码/测试状态，不推断用户价值"],
        ["用户问题/目标/约束", "user_confirmed Canvas", "回到 Guide 让用户确认"],
        ["模拟访谈场景", "simulated_research", "必须在同一主张局部说明“模拟/人工构造”"],
    ], widths=[4.2, 6.0, 6.3])
    add_heading(doc, "13.3 Validator", 2)
    add_bullets(doc, [
        "检查 required headings、citation existence、project scope 和 Canvas consistency。",
        "检查 source_type disclosure：simulated/model_hypothesis 不能只在文档开头统一免责声明。",
        "检查 source-backed claim 是否有证据 link，link 的 source/chunk 是否属于当前项目。",
        "检查 expected source type 与实际类型是否冲突。",
        "发现 unresolved 或 claim mismatch 时阻止 approval，或明确转人工。",
        "最多两轮 repair；达到预算、无新证据或冲突未解时终止。",
    ])
    add_callout(doc, "方法边界", "Claim-Evidence Ledger 提高追溯与错误暴露能力，但不能仅靠词法匹配证明一句话在语义上完全被来源支持。正式业务使用仍需要人工审查或另行验证的 NLI/LLM Judge 评测，且 Judge 不能复用测试集指导参数选择。", AMBER_SOFT, AMBER)

    add_heading(doc, "14. 数据模型、API 与审计", 1)
    add_heading(doc, "14.1 数据模型", 2)
    add_table(doc, ["表", "关键字段", "用途"], [
        ["projects", "id,title,summary,status", "项目身份"],
        ["project_canvas / versions", "problem,target_users,goals,metrics,constraints,version", "当前/不可变项目合同"],
        ["guided_sessions", "current_step,state_json,status", "Coach 当前状态"],
        ["guided_messages", "role,type,content,metadata,time", "澄清对话回放"],
        ["project_decisions", "options,selected,rationale,status", "方案选择记录"],
        ["sources", "type,authority,basis,url,publisher,time,status,sha", "原始来源与 provenance"],
        ["source_chunks", "source_id,project_id,index,content", "确定性切片与项目边界"],
        ["retrieval_runs / hits", "query,profile,weights,candidates,scores", "检索可回放"],
        ["documents / versions", "doc_type,version,canvas_version,status,content,citations", "不可变文档"],
        ["document_claims", "claim_type,support_status,section,text", "结构化主张"],
        ["claim_evidence_links", "source,chunk,relation,retrieval_run", "主张与证据关系"],
        ["generation_runs / issues", "rounds,terminal_state,issue", "bounded Loop 证据"],
        ["handoff_runs", "target,status,manifest,sha", "交接运行与包身份"],
        ["audit_events", "actor,action,entity,payload,time", "统一审计"],
    ], widths=[4.2, 7.0, 5.3], font_size=7.2)
    add_heading(doc, "14.2 API", 2)
    add_table(doc, ["域", "主要接口"], [
        ["Projects/Canvas", "GET/POST /api/projects；GET/PUT /api/projects/{id}/canvas"],
        ["Guide", "GET /guide；POST /guide/respond；POST /guide/apply；POST /guide/reset"],
        ["Sources", "GET /api/source-guidance；GET/POST /sources；POST /sources/guided；POST /sources/upload"],
        ["Retrieval", "GET /api/retrieval/profiles；POST /retrieve；GET /retrieval-runs；GET /api/retrieval/runs/{run_id}"],
        ["Documents", "POST /generate；GET /api/documents/{version_id}；GET /claims；POST /validate；POST /approve；GET /export"],
        ["Handoff", "GET /handoff/readiness；POST /handoff/export"],
        ["Operations", "GET /api/health；GET /api/audit；GET /api/tools"],
    ], widths=[3.4, 13.1], font_size=7.4)
    add_heading(doc, "14.3 迁移与兼容", 2)
    add_bullets(doc, [
        "SQLite 使用 additive/idempotent migration；旧 v1.1.0 projects/sources 不删除。",
        "旧 authority 数字若无依据，标记 legacy/manual numeric score without recorded basis。",
        "现有 API 保持兼容；list-returning retrieval service 保留为内部 legacy wrapper。",
        "新 UI 调用完整 trace API，不再只拿 items。",
    ])

    add_heading(doc, "15. 真实 AI Coding 交接 ZIP", 1)
    add_callout(doc, "产品判断", "AI Coding 交接不能只是一段 Prompt 或“已提供 MCP 接口”的描述。v2 将交接定义为一个有 readiness Gate、有批准版本、有文件清单和哈希、可由外部 Agent 消费的真实 ZIP。", MINT_SOFT, MINT)
    add_heading(doc, "15.1 Readiness Gate", 2)
    add_table(doc, ["条件", "通过标准", "失败行为"], [
        ["Canvas", "存在已确认版本", "BLOCKED：回到 Guide/Advanced 保存"],
        ["PRD", "approved + validation passed", "BLOCKED：生成、验证、人工阅读并批准"],
        ["TechDoc", "approved + validation passed", "BLOCKED：同上"],
        ["未解决主张", "允许存在但必须进入包并显示数量", "不得静默删除或包装成事实"],
    ], widths=[3.3, 7.0, 6.2])
    add_heading(doc, "15.2 包内容", 2)
    add_table(doc, ["文件", "作用"], [
        ["README_FIRST.md", "阅读顺序、版本和边界"],
        ["APPROVED_CONTEXT.md", "Canvas 问题、用户、目标、非目标、指标和约束"],
        ["PRD_APPROVED.md / TECHDOC_APPROVED.md", "批准且验证通过的不可变文档"],
        ["CLAIM_LEDGER.json", "所有主张、状态、解释和证据 link"],
        ["SOURCE_MANIFEST.json", "来源类型、URL、发布者、authority basis、SHA"],
        ["RETRIEVAL_TRACE.json", "文档相关检索运行和 hits"],
        ["ACCEPTANCE_TESTS.md", "由成功指标、约束和安全规则形成的验收"],
        ["IMPLEMENTATION_TASKS.json", "按目标拆解的任务、依赖和完成定义"],
        ["AGENTS.md", "AI Coding 不得越过的版本、非目标、审批和测试规则"],
        ["HANDOFF_MANIFEST.json", "每个文件 SHA-256、项目/Canvas/文档版本、run ID、target client"],
    ], widths=[5.3, 11.2])
    add_heading(doc, "15.3 交接验收", 2)
    add_bullets(doc, [
        "Zip 中每个 manifest 记录的文件都存在且哈希一致。",
        "HTTP 响应返回 X-Handoff-SHA256 和 X-Handoff-Run-ID。",
        "handoff_runs 保存 target client、status、manifest、package SHA 和 time。",
        "Audit 记录 handoff_package_exported。",
        "Codex/Claude Code/Cursor/Generic 只是消费目标标签，不改变核心包语义。",
    ])

    add_heading(doc, "16. 指标体系、RAG Evals 与实验协议", 1)
    add_heading(doc, "16.1 指标原则", 2)
    add_callout(doc, "指标边界", "“生成更快、文本更长、界面更像 AI”不是成功。必须同时观察完成率、认知负担、来源理解、主张支持、人工修改、交接可执行性和安全错误。", BLUE_SOFT, BLUE)
    add_table(doc, ["层级", "指标", "定义/目的"], [
        ["北极星", "Evidence-backed Project Completion Rate", "完成 Idea→Canvas→Evidence→Approved Docs→Handoff，且来源/主张状态合格的项目比例"],
        ["漏斗", "Project → Guide → Canvas → Source → Draft → Approved → Handoff", "定位中断阶段"],
        ["新手可用性", "Guide completion / time / help requests / backtracks", "验证引导是否降低而非增加负担"],
        ["来源理解", "Source Type Accuracy / Boundary Comprehension", "用户能否正确判断资料能/不能证明什么"],
        ["检索", "Recall@K / MRR / Candidate & Returned Count", "找到预期证据与阅读负担"],
        ["引用", "Citation Precision / Claim Support Rate", "引用是否支持对应主张"],
        ["安全", "Unsupported Claim Rate / Disclosure Accuracy / Cross-project leakage", "防止无证据、错误升级和越界"],
        ["文档", "Rubric pass / edit distance / issue count / rounds", "质量提升是否只是文本更长"],
        ["Handoff", "Task executability / requirement deviation / rework", "交接包是否真正减少偏离和返工"],
    ], widths=[2.6, 5.7, 8.2], font_size=7.4)
    add_heading(doc, "16.2 RAG 评测 UI", 2)
    add_table(doc, ["新手看到", "后台指标", "解释"], [
        ["系统有没有找到应该找到的资料", "Evidence Recall@K", "Golden query 是否召回标注 chunk"],
        ["引用是否真的支持这句话", "Citation Precision / Claim Support Rate", "不是只检查 citation ID 存在"],
        ["有没有把无证据内容写成事实", "Unsupported Claim Rate", "unresolved/model suggestion 是否正确披露"],
        ["当前档位为什么这样设置", "Profile, Top-K, weights, validation status", "显示依据和 trade-off"],
    ], widths=[5.0, 5.2, 6.3])
    add_heading(doc, "16.3 实验协议", 2)
    add_numbered(doc, [
        "建立冻结 query/evidence 数据集，记录来源版本和 SHA-256。",
        "仅在 validation/select split 比较 quick/balanced/thorough 或新配置。",
        "同时报告 Recall@K、Citation Precision、Unsupported Claim Rate、候选数、延迟和人工阅读量。",
        "模型/参数锁定后只运行一次 frozen report；不得根据 frozen 结果继续调参。",
        "按来源类型、任务阶段和难例分组；保留失败案例。",
        "RAG 没有稳定净收益时退回 BM25/更简单基线，不为技术复杂度保留。",
    ])
    add_table(doc, ["实验", "对照", "成功标准", "失败退出条件"], [
        ["Guided vs blank Canvas", "同一 idea 直接填表 vs 引导", "完成率↑、帮助次数↓、时间不显著恶化、合同质量↑", "用户认为只是多一步，且无质量收益"],
        ["Evidence Panel", "仅最终引用 vs 运行/主张追溯", "来源判断准确率↑、不信任/回查成本↓", "用户不使用面板或理解率无改善"],
        ["Retrieval Profiles", "BM25 / current hybrid / candidate", "validation 上 Recall/Precision 有稳定净收益", "仅在 test 偶然提升或阅读负担显著增加"],
        ["Handoff ZIP", "普通 PRD 粘贴 vs versioned package", "任务偏差/返工↓，验收通过率↑", "AI Agent 不使用 manifest 或无实际改善"],
    ], widths=[3.6, 4.3, 4.8, 3.8], font_size=7.2)

    add_heading(doc, "17. 用户测试、上线策略与运营埋点", 1)
    add_heading(doc, "17.1 目标用户任务测试", 2)
    add_table(doc, ["阶段", "样本", "任务", "观察"], [
        ["Alpha", "本人 + 2–3 名非目标人员", "从 rough idea 到 Canvas 与来源", "发现交互/术语/运行缺陷，不做市场结论"],
        ["Closed Beta", "8–12 名目标用户：求职者、转岗 PM、独立开发者", "Guide、证据、检索、文档、交接全流程", "完成率、求助、回退、理解、修改量、信任校准"],
        ["RAG Eval", "冻结项目与 golden query/evidence", "多个 profile/基线离线比较", "validation-only 选择、frozen report"],
        ["Handoff Pilot", "2–4 个真实可编码小项目", "Codex/Claude Code/Cursor 消费 ZIP", "偏差、返工、测试通过、边界违反"],
    ], widths=[3.0, 4.4, 5.1, 4.0])
    add_heading(doc, "17.2 关键埋点", 2)
    add_table(doc, ["事件", "关键属性"], [
        ["project_created", "project_id, entry_mode"],
        ["guide_answered", "step, answer_type, used_example, duration"],
        ["guide_canvas_applied", "canvas_version, confirmed_fields, selected_solution"],
        ["source_added", "origin_kind, source_type, needs_confirmation, url_present"],
        ["retrieval_run_created", "profile, top_k_source, candidates, returned, purpose"],
        ["claim_viewed", "claim_type, evidence_count, source_opened"],
        ["document_validated", "doc_type, issue_codes, rounds"],
        ["document_approved", "doc_type, version, actor"],
        ["handoff_exported", "target_client, manifest_files, unresolved_count, sha"],
    ], widths=[5.0, 11.5])
    add_heading(doc, "17.3 上线 Gate", 2)
    add_bullets(doc, [
        "所有自动测试、compileall、JS syntax 和 fresh-package rerun 通过。",
        "默认 Guided Mode 在桌面/移动截图检查无重叠、截断或不可见切换。",
        "DOCX PRD 经 render → page PNG 全页检查。",
        "迁移旧 1.1.0 DB 不丢 project/source。",
        "交接 ZIP manifest 重新计算全部文件哈希一致。",
        "README/PRD/CLAIM_BOUNDARY 与实际代码能力一致。",
    ])

    add_heading(doc, "18. 风险、失败退出条件与主张边界", 1)
    add_table(doc, ["风险", "后果", "当前控制", "失败退出条件"], [
        ["引导过长", "小白更快流失", "一次一问、进度、选择/示例、允许 Advanced", "完成率不升且耗时/回退显著增加 → 减少步骤"],
        ["来源分类仍难", "错误证据升级", "plain-language category + allowed/limitations + confirmation", "错误分类持续高 → 改为更少类别或人工确认"],
        ["RAG 复杂但无收益", "参数和维护负担增加", "命名 baseline、Run、冻结评测协议", "无稳定净收益 → 简化为 BM25/手选"],
        ["Claim 对齐假安全", "有链接但语义不支持", "source-type + lexical fit + unresolved + human review", "错误支持率高 → 禁止自动事实性表述"],
        ["Handoff 无人使用", "只是作品集装饰", "真实 ZIP + target pilot + task metrics", "Agent 不读取/无返工收益 → 收缩文件集"],
        ["Agent 越权", "错误写入/审批/发布", "ToolRegistry 风险级、host confirmation、缺危险工具", "发现绕过 Gate → 禁用 LLM path"],
        ["单次反馈被夸大", "论文/简历/产品主张失真", "反馈性质明确写入 PRD/Claim Boundary", "任何外部材料写“已市场验证” → 立即修正"],
    ], widths=[3.5, 3.6, 5.5, 4.0], font_size=7.0)
    add_heading(doc, "18.1 可表述 / 不可表述", 2)
    add_table(doc, ["可以表述", "不能表述"], [
        ["完成可运行的 novice-first Guided Evidence Workspace。", "已由市场验证或获得真实留存/付费结果。"],
        ["实现受限 PM Coach、状态持久化和显式 Canvas Gate。", "Agent 可自主替代产品经理。"],
        ["实现可回放 retrieval trace 和 structured Claim-Evidence links。", "系统自动保证每个引用都语义正确。"],
        ["实现真实 AI Coding 交接 ZIP 与 SHA-256 manifest。", "已完成生产级 Codex/Claude/Cursor 集成闭环。"],
        ["实现本地 MCP 资源和工具。", "已完成远程 MCP OAuth、企业权限和多人协作。"],
        ["当前 RAG 配置是可审计 baseline。", "Top-K=8 / weights 是最佳参数。"],
    ], widths=[8.2, 8.3])

    add_heading(doc, "19. Roadmap 与当前实现审计", 1)
    add_table(doc, ["阶段", "目标", "当前状态", "进入/退出 Gate"], [
        ["M0 v1.1 Evidence Workspace", "项目/Canvas/来源/RAG/文档/Loop/MCP 可运行", "历史完成", "被 v2 新手形态替代"],
        ["M1 v2 Guided Evidence Workspace", "默认引导、双模式、Trace、Claims、Handoff", "已实现", "自动测试 + UI/DOCX/包验证"],
        ["M2 Real User Validation", "验证新手价值和认知负担", "待执行", "8–12 名目标用户；失败则收缩流程"],
        ["M3 Frozen RAG Eval", "验证 profile/权重/Top-K", "待执行", "validation-only 选择 + frozen report"],
        ["M4 Handoff Pilot", "验证 AI Coding 实际消费", "待执行", "2–4 个真实项目，偏差/返工指标"],
        ["M5 Conditional Integrations", "网页抓取、远程 MCP、团队协作", "条件式", "必须有真实需求、安全与权限设计"],
    ], widths=[4.0, 5.2, 2.8, 4.5])
    add_heading(doc, "19.1 当前实现证据", 2)
    add_bullets(doc, [
        "FastAPI + SQLite 本地应用；旧数据库 additive migration。",
        "Guided/Advanced 双模式 UI，桌面和移动静态截图无运行错误。",
        "GuidedProjectService、SourceGuidanceService、ProjectRetrievalService、ClaimLedgerService、HandoffService。",
        "named retrieval profiles、run/hit 持久化、project-scoped ranking。",
        "文档生成不再按索引轮换引用；structured claims 持久化。",
        "真实 Handoff ZIP、file/package SHA-256、audit、MCP wrappers。",
        "截至本文档生成时，完整仓库 77 项 pytest 自动测试通过；最终 fresh-package 结果以 VERIFICATION_REPORT.md 为准。",
    ])
    add_heading(doc, "19.2 本轮未完成", 2)
    add_bullets(doc, [
        "真实目标用户任务测试与统计结果。",
        "冻结 RAG 查询/证据评测集和最佳参数选择。",
        "自动网页抓取、内容失效检测与合规治理。",
        "生产远程 MCP、SSO、多租户与实时协作。",
        "真实在线 Agent/Coding 平台回写闭环。",
        "大规模性能、成本、可用性和安全压测。",
    ])
    add_callout(doc, "最终产品结论", "InsightForge 2.0 已从“技术能力展示”推进为“围绕新手完成真实任务的可运行产品设计”。但这只能证明实现与设计迭代已完成，不能证明市场价值成立。下一阶段最重要的工作不是继续堆 Agent/模型，而是执行真实用户任务测试、冻结 RAG 评测和 Handoff 实际消费验证。", MINT_SOFT, MINT)

    add_heading(doc, "附录 A：验收清单", 1)
    add_table(doc, ["域", "验收项", "结果要求"], [
        ["Guided", "默认模式、六步 rail、one-question、why/examples/choices、apply Gate", "全部可用"],
        ["Source", "普通分类、allowed/limitations、URL/Publisher/Time/Basis/SHA", "可保存和显示"],
        ["RAG", "profiles、run、hits、scores、project scope、空结果", "可回放且无跨项目"],
        ["Claims", "四类状态、source/chunk/run links、type mismatch validation", "可查询和审计"],
        ["Documents", "immutable version、bounded Loop、approval、DOCX export", "Gate 生效"],
        ["Handoff", "readiness、文件清单、hash、run、audit", "fresh unzip 校验一致"],
        ["UI", "desktop/mobile layout、mode labels、no page errors", "人工截图通过"],
        ["PRD", "rendered pages no clipping/overlap", "全页 PNG 检查通过"],
        ["Package", "clean ZIP、fresh extraction、full tests", "0 failures"],
    ], widths=[2.8, 8.3, 5.4])

    add_heading(doc, "附录 B：参考与证据来源", 1)
    add_bullets(doc, [
        "InsightForge_完整产品需求与策略文档_v1.0（2026-08-24）：产品策略、用户、竞品、RAG 与版本规划基线。",
        "InsightForge v1.1.0 可运行代码包：Canvas、来源、检索、文档 Loop、Function Calling、MCP 与测试基线。",
        "2026-08-26 一次定性产品经理反馈：UI、新手认知负担、RAG 配置解释、Handoff 与黑箱问题。",
        "InsightForge 2.0 代码与自动测试：Guided/Advanced、provenance、retrieval trace、claims、handoff。",
        "本轮没有将模拟研究、代码实现或单次反馈表述为真实市场验证。",
    ])

    return doc


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.core_properties.title = "InsightForge 2.0 完整产品需求与策略文档"
    doc.core_properties.subject = "Guided Evidence Workspace"
    doc.core_properties.author = "InsightForge Product & Engineering"
    doc.core_properties.keywords = "AI产品经理, PRD, RAG, Agent, MCP, 可追溯, AI Coding Handoff"
    doc.core_properties.comments = "Implementation-aligned v2.0.0; qualitative feedback is not market validation."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
