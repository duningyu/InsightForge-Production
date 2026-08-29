from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt


class ArtifactExporter:
    """Serialize immutable document-version records into portable artifacts."""

    @staticmethod
    def _normalized(version: dict[str, Any]) -> dict[str, Any]:
        result = dict(version)
        if "citations" not in result and "citations_json" in result:
            result["citations"] = json.loads(result.pop("citations_json"))
        return result

    def to_markdown(self, version: dict[str, Any]) -> str:
        normalized = self._normalized(version)
        content = str(normalized.get("content", "")).strip()
        if not content:
            raise ValueError("document version contains no content")
        return content + "\n"

    def to_json_text(self, version: dict[str, Any]) -> str:
        return json.dumps(self._normalized(version), ensure_ascii=False, indent=2)

    def to_json_bytes(self, version: dict[str, Any]) -> bytes:
        return self.to_json_text(version).encode("utf-8")

    def to_docx_bytes(self, version: dict[str, Any]) -> bytes:
        normalized = self._normalized(version)
        content = str(normalized.get("content", "")).strip()
        if not content:
            raise ValueError("document version contains no content")

        document = Document()
        properties = document.core_properties
        properties.title = f"{normalized.get('doc_type', 'document')} v{normalized.get('version', '')}".strip()
        properties.subject = "InsightForge cited artifact"
        properties.comments = (
            f"version_id={normalized.get('id', '')}; "
            f"project_id={normalized.get('project_id', '')}; "
            f"canvas_version={normalized.get('canvas_version', '')}"
        )

        in_code = False
        code_lines: list[str] = []
        for raw_line in content.splitlines():
            line = raw_line.rstrip()
            if line.strip().startswith("```"):
                if in_code:
                    paragraph = document.add_paragraph()
                    run = paragraph.add_run("\n".join(code_lines))
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)
                    code_lines = []
                    in_code = False
                else:
                    in_code = True
                continue
            if in_code:
                code_lines.append(line)
                continue
            if not line.strip():
                document.add_paragraph("")
                continue
            heading = re.match(r"^(#{1,3})\s+(.+)$", line)
            if heading:
                document.add_heading(heading.group(2).strip(), level=len(heading.group(1)))
                continue
            if line.startswith("- "):
                document.add_paragraph(line[2:].strip(), style="List Bullet")
                continue
            if re.match(r"^\d+\.\s+", line):
                document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
                continue
            quote = line.startswith("> ")
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line[2:] if quote else line)
            if quote:
                run.italic = True

        if code_lines:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = "Courier New"
            run.font.size = Pt(9)

        document.add_page_break()
        document.add_heading("导出元数据", level=1)
        metadata = [
            ("Version ID", normalized.get("id")),
            ("Project ID", normalized.get("project_id")),
            ("Document Type", normalized.get("doc_type")),
            ("Version", normalized.get("version")),
            ("Canvas Version", normalized.get("canvas_version")),
            ("Validation Status", normalized.get("validation_status")),
            ("Approval Status", normalized.get("status")),
        ]
        for key, value in metadata:
            document.add_paragraph(f"{key}: {value}")
        citations = normalized.get("citations") or []
        document.add_heading("引用清单", level=2)
        for citation in citations:
            document.add_paragraph(str(citation), style="List Bullet")

        output = BytesIO()
        document.save(output)
        return output.getvalue()
