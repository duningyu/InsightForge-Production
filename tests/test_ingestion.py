import io
import json

from docx import Document

from app.ingestion import chunk_text, extract_text
from app.services.sources import SourceService


def test_chunk_text_is_deterministic_and_overlapping():
    text = "第一段。第二段。第三段。" * 40
    chunks = chunk_text(text, max_chars=120, overlap=20)
    assert len(chunks) > 1
    assert chunks == chunk_text(text, 120, 20)
    assert all(len(chunk) <= 120 for chunk in chunks)


def test_extract_text_supports_json_and_docx():
    json_bytes = json.dumps({"title": "需求", "items": ["引用", "版本"]}, ensure_ascii=False).encode()
    assert "引用" in extract_text("brief.json", json_bytes)

    document = Document()
    document.add_heading("项目证据", level=1)
    document.add_paragraph("需要保留 source_id。")
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = extract_text("evidence.docx", buffer.getvalue())
    assert "source_id" in extracted


def test_source_type_is_preserved(db):
    source = SourceService(db).add_source(
        project_id="project_insightforge_demo",
        title="模拟访谈",
        source_type="simulated_research",
        authority=0.4,
        content="这是人工构造的访谈样本。",
        filename="sample.txt",
    )
    assert source["source_type"] == "simulated_research"
    chunks = db.fetch_all("SELECT * FROM source_chunks WHERE source_id = ?", (source["id"],))
    assert chunks


def test_invalid_source_type_is_rejected(db):
    service = SourceService(db)
    try:
        service.add_source(
            project_id="project_insightforge_demo",
            title="错误来源",
            source_type="real_research_without_contract",
            authority=0.8,
            content="bad",
            filename="bad.txt",
        )
    except ValueError as exc:
        assert "invalid source_type" in str(exc)
    else:
        raise AssertionError("invalid source type should fail")
