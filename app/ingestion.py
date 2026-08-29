from __future__ import annotations

import io
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.db import chunk_text

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".tsv", ".yaml", ".yml"}
MAX_UPLOAD_BYTES = 12 * 1024 * 1024


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("unable to decode text as UTF-8 or GB18030")


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"file exceeds {MAX_UPLOAD_BYTES} bytes")
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        text = _decode_text(data)
    elif suffix == ".json":
        parsed = json.loads(_decode_text(data))
        text = json.dumps(parsed, ensure_ascii=False, indent=2)
    elif suffix == ".docx":
        document = Document(io.BytesIO(data))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "").strip() for page in reader.pages)
    else:
        raise ValueError(f"unsupported file type: {suffix or 'no extension'}")
    normalized = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").split("\n")).strip()
    if not normalized:
        raise ValueError("source contains no extractable text")
    return normalized


__all__ = ["chunk_text", "extract_text", "MAX_UPLOAD_BYTES"]
